from fastapi import Request
from fastapi.templating import Jinja2Templates
from fastapi.responses import StreamingResponse, JSONResponse
from config import BASE_URL
from helpers.database_helper import (
    fetchrecords,
    main_query,
    get_dcr_summary,
    delete_dcr_report,
    clear_all_dcr_reports,
)
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from io import BytesIO

templates = Jinja2Templates(directory="templates")

async def extracterindex(request: Request):
    result = await fetchrecords(main_query())
    summary = get_dcr_summary(result)
    context = {
        "BASE_URL": BASE_URL,
        "reports": result,
        "summary": summary
    }
    return templates.TemplateResponse(request=request, name="extracterindex.html", context=context)

async def delete_record(request: Request, record_id: str):
    success = await delete_dcr_report(record_id)
    return JSONResponse({
        "status": success,
        "message": "Record deleted successfully." if success else "Failed to delete record."
    })

async def clear_all_records(request: Request):
    count = await clear_all_dcr_reports()
    return JSONResponse({
        "status": True,
        "deleted_count": count,
        "message": f"{count} records cleared successfully."
    })

async def downloadExcel(request: Request):

    result = await fetchrecords(main_query())

    wb = Workbook()
    ws = wb.active
    ws.title = "DCR REPORT"

    # ------------------------
    # Header
    # ------------------------

    ws.merge_cells("A1:A2")
    ws.merge_cells("B1:B2")

    ws["A1"] = "FILE NAME"
    ws["B1"] = "CINEMA NAME"

    header_groups = [
        ("FRIDAY", "C", "E"),
        ("SATURDAY", "F", "H"),
        ("SUNDAY", "I", "K"),
        ("MONDAY", "L", "N"),
        ("TUESDAY", "O", "Q"),
        ("WEDNESDAY", "R", "T"),
        ("THURSDAY", "U", "W"),
        ("TOTAL", "X", "Z"),
    ]

    for title, start, end in header_groups:
        ws.merge_cells(f"{start}1:{end}1")
        ws[f"{start}1"] = title

    headers = [
        "SHOW","ADMITS","NET",
        "SHOW","ADMITS","NET",
        "SHOW","ADMITS","NET",
        "SHOW","ADMITS","NET",
        "SHOW","ADMITS","NET",
        "SHOW","ADMITS","NET",
        "SHOW","ADMITS","NET",
        "NET SHOW","NET ADM","GRAND TOTAL NET"
    ]

    col = 3

    for h in headers:
        ws.cell(row=2, column=col).value = h
        col += 1

    # ------------------------
    # Data Rows
    # ------------------------

    for row in result:

        ws.append([

            row.get("file_name",""),
            row.get("cinema_name",""),

            row.get("fri_show",0),
            row.get("fri_admits",0),
            row.get("fri_net",0),

            row.get("sat_show",0),
            row.get("sat_admits",0),
            row.get("sat_net",0),

            row.get("sun_show",0),
            row.get("sun_admits",0),
            row.get("sun_net",0),

            row.get("mon_show",0),
            row.get("mon_admits",0),
            row.get("mon_net",0),

            row.get("tue_show",0),
            row.get("tue_admits",0),
            row.get("tue_net",0),

            row.get("wed_show",0),
            row.get("wed_admits",0),
            row.get("wed_net",0),

            row.get("thu_show",0),
            row.get("thu_admits",0),
            row.get("thu_net",0),

            row.get("net_show", 0),
            row.get("net_adm", 0),
            row.get("grand_total", 0)

        ])

    num_records = len(result)
    total_row_idx = num_records + 3

    # Add Grand Total Summary Row if there are records
    if num_records > 0:
        ws.cell(row=total_row_idx, column=1).value = "GRAND TOTAL"
        ws.cell(row=total_row_idx, column=2).value = f"{num_records} Cinemas"
        for col_idx in range(3, 27):
            col_letter = get_column_letter(col_idx)
            ws.cell(row=total_row_idx, column=col_idx).value = f"=SUM({col_letter}3:{col_letter}{total_row_idx - 1})"

    # ------------------------
    # Styling, Borders & Formats
    # ------------------------

    thin = Side(style="thin")
    double_bottom = Side(style="double")
    border_standard = Border(left=thin, right=thin, top=thin, bottom=thin)
    border_total_row = Border(left=thin, right=thin, top=thin, bottom=double_bottom)

    header_fill_days = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    header_fill_total = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")
    total_row_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")

    net_cols = {5, 8, 11, 14, 17, 20, 23, 26}   # E, H, K, N, Q, T, W, Z
    adm_cols = {4, 7, 10, 13, 16, 19, 22, 25}   # D, G, J, M, P, S, V, Y
    show_cols = {3, 6, 9, 12, 15, 18, 21, 24}   # C, F, I, L, O, R, U, X

    for r_idx, row in enumerate(ws.iter_rows(), start=1):
        is_header = (r_idx in [1, 2])
        is_total_summary = (num_records > 0 and r_idx == total_row_idx)

        for c_idx, cell in enumerate(row, start=1):
            if is_header:
                cell.border = border_standard
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.fill = header_fill_total if c_idx >= 24 else header_fill_days
            elif is_total_summary:
                cell.border = border_total_row
                cell.font = Font(bold=True)
                cell.fill = total_row_fill
                if c_idx in net_cols:
                    cell.number_format = '#,##0.00'
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif c_idx in adm_cols or c_idx in show_cols:
                    cell.number_format = '#,##0'
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.border = border_standard
                if c_idx in net_cols:
                    cell.number_format = '#,##0.00'
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif c_idx in adm_cols:
                    cell.number_format = '#,##0'
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif c_idx in show_cols:
                    cell.number_format = '#,##0'
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

    # Row Heights
    ws.row_dimensions[1].height = 25
    ws.row_dimensions[2].height = 24
    if num_records > 0:
        ws.row_dimensions[total_row_idx].height = 24

    # ------------------------
    # Column Widths
    # ------------------------

    for column in ws.columns:
        max_length = 0
        col_letter = get_column_letter(column[0].column)
        for cell in column:
            try:
                if cell.value and not str(cell.value).startswith("="):
                    max_length = max(max_length, len(str(cell.value)))
            except:
                pass
        ws.column_dimensions[col_letter].width = max(max_length + 3, 10)

    # Explicit padding for text & total columns
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 26)
    ws.column_dimensions["B"].width = max(ws.column_dimensions["B"].width or 0, 30)
    ws.column_dimensions["X"].width = max(ws.column_dimensions["X"].width or 0, 13)
    ws.column_dimensions["Y"].width = max(ws.column_dimensions["Y"].width or 0, 13)
    ws.column_dimensions["Z"].width = max(ws.column_dimensions["Z"].width or 0, 20)

    # Freeze Header
    ws.freeze_panes = "C3"

    # ------------------------
    # Download
    # ------------------------

    stream = BytesIO()

    wb.save(stream)

    stream.seek(0)

    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition":"attachment; filename=DCR_Report.xlsx"
        }
    )